#!/usr/bin/env python3
"""Compare every FDI story artifact with its frozen CSV/XLSX source where available.

This is a source-to-artifact check, not a prose or conceptual accuracy certificate.
Exit nonzero on any cell mismatch or an artifact without a comparison rule.
"""
import csv
import json
import math
from hashlib import sha256
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
RAW = ROOT / "data/raw/unctad-wir"
MANIFEST = ROOT / "data/catalog/unctad-wir-fdi-manifest.json"
rows = list(csv.DictReader((RAW / "unctadstat/US_FdiFlowsStock.csv").open()))
by_key = {(int(r["Year"]), r["Economy Label"], r["Flow Label"], r["Direction Label"]): r for r in rows}
population = {r["Economy Label"]: float(r["Absolute value in thousands"]) * 1000
              for r in csv.DictReader((RAW / "unctadstat/US_PopTotal.csv").open())
              if r["Year"] == "2025" and r["Absolute value in thousands"]}
wb_payload = json.loads((RAW / "worldbank/BX_KLT_DINV_WD_GD_ZS.json").read_text())
wb = {(r["countryiso3code"], int(r["date"])): r["value"] for r in wb_payload[1] if r["value"] is not None}
MONEY = "Millions of US$ at current prices"
GFCF = "Percentage of gross Fixed Capital Formation"
GDP = "Percentage of gross Domestic Product"
errors = []
checked = 0
unverified = []
for source_file in json.loads((RAW / "manifest.json").read_text())["files"]:
    file_path = RAW / source_file["path"]
    if not file_path.exists():
        errors.append({"missing_source": source_file["path"]})
    elif sha256(file_path.read_bytes()).hexdigest() != source_file["sha256"]:
        errors.append({"changed_source": source_file["path"]})


def jsround(v, decimals=0):
    scale = 10**decimals
    return math.floor(v * scale + 0.5) / scale


def raw(year, economy, field=MONEY, kind="Flow", direction="Inward"):
    item = by_key.get((year, economy, kind, direction))
    return float(item[field]) if item and item[field] else None


def same(label, got, expected):
    global checked
    checked += 1
    if isinstance(got, (int, float)) and isinstance(expected, (int, float)):
        equal = math.isclose(got, expected, abs_tol=1e-8)
    else:
        equal = got == expected
    if not equal:
        errors.append({"cell": label, "artifact": got, "source": expected})


def workbook(num):
    return load_workbook(RAW / f"wir2026/annex/wir26_tab{num:02d}.xlsx", read_only=True, data_only=True).active


def annual_annex(num, economy):
    sheet = workbook(num)
    years = [c.value for c in sheet[3]][1:]
    for row in sheet.iter_rows(min_row=4):
        if str(row[0].value).strip() == economy:
            return {int(y): float(row[i + 1].value) for i, y in enumerate(years)
                    if str(y).isdigit() and isinstance(row[i + 1].value, (int, float))}
    raise ValueError(f"India/economy row absent in annex {num}: {economy}")


def mne_counts(num):
    sheet = workbook(num)
    data = list(sheet.values)
    cols = [i for row in data for i, value in enumerate(row) if value == "Home economy"]
    if not cols:
        raise ValueError(f"Home economy column absent in annex {num}")
    col = cols[0]
    return Counter(str(row[col]).strip() for row in data[4:] if row[col] not in (None, "", "Home economy"))


gfcf = {"IND": "India", "VNM": "Viet Nam", "CHN": "China", "KOR": "Republic of Korea",
        "POL": "Poland", "PAK": "Pakistan", "BGD": "Bangladesh", "LKA": "Sri Lanka"}
flow = {"IND": "India", "CHN": "China", "SEA": "Developing economies: South-eastern Asia"}
stock_peers = ["Viet Nam", "Thailand", "Malaysia", "Brazil", "Poland", "Mexico",
               "Indonesia", "China", "India", "Bangladesh"]
caribbean_financial_centres = {
    "Anguilla", "Antigua and Barbuda", "Aruba", "Bahamas", "Barbados",
    "British Virgin Islands", "Cayman Islands", "Curaçao", "Dominica",
    "Grenada", "Montserrat", "Saint Kitts and Nevis", "Saint Lucia",
    "Saint Vincent and the Grenadines", "Sint Maarten", "Turks and Caicos Islands",
}
ind_in = {y: raw(y, "India") for y in range(1990, 2026)}
ind_out = {y: raw(y, "India", direction="Outward") for y in range(1990, 2026)}
dev_total = {y: raw(y, "Developing economies") for y in range(1990, 2026)}
annex_in = annual_annex(1, "India")
for year in range(1990, 2026):
    same(f"UNCTAD CSV vs WIR annex table 01:India:{year}", ind_in[year], jsround(annex_in[year], 3))
announced_project_counts = annual_annex(17, "India")
same("WIR annex table 17:India:2024 announced projects", announced_project_counts[2024], 1089)
same("WIR annex table 17:India:2025 announced projects", announced_project_counts[2025], 1037)

for entry in json.loads(MANIFEST.read_text()):
    if entry.get("status") != "ready":
        continue
    artifact = json.loads((ROOT / entry["artifact"]).read_text())
    aid = artifact["indicatorId"]
    observations = artifact.get("observations", [])
    table_rows = artifact.get("rows", [])
    if aid.startswith("extfin.fdi.dev.gfcf_share."):
        econ = gfcf[aid.split(".")[-2]]
        for o in observations:
            year = int(o["date"][:4])
            same(f"{aid}:{year}", o["value"], jsround(raw(year, econ, GFCF), 2))
    elif aid.startswith("extfin.fdi.dev.gdp_share_wb."):
        code = aid.split(".")[-2]
        for o in observations:
            year = int(o["date"][:4])
            same(f"{aid}:{year}", o["value"], jsround(wb[code, year], 3))
    elif aid.startswith("extfin.fdi.dev.inward_flow."):
        econ = flow[aid.split(".")[-2]]
        for o in observations:
            year = int(o["date"][:4])
            same(f"{aid}:{year}", o["value"], jsround(raw(year, econ), 2))
    elif aid in ("extfin.fdi.dev.inward_stock.IN.usd", "extfin.fdi.dev.outward_stock.IN.usd"):
        direction = "Inward" if "inward" in aid else "Outward"
        for o in observations:
            year = int(o["date"][:4])
            same(f"{aid}:{year}", o["value"], jsround(raw(year, "India", kind="Stock", direction=direction), 2))
    elif aid == "extfin.fdi.dev.world_rank.IN":
        for o in observations:
            year = int(o["date"][:4])
            eligible = sorted((float(r[MONEY]), r["Economy Label"]) for r in rows
                              if r["Year"] == str(year) and r["Flow Label"] == "Flow"
                              and r["Direction Label"] == "Inward" and r[MONEY]
                              and r["Economy"].isdigit() and len(r["Economy"]) <= 3
                              and r["Economy Label"] not in caribbean_financial_centres)
            expected = next(i + 1 for i, (_, e) in enumerate(reversed(eligible)) if e == "India")
            same(f"{aid}:{year}", o["value"], expected)
    elif aid == "extfin.fdi.dev.india_share_developing.pct":
        for o in observations:
            year = int(o["date"][:4])
            same(f"{aid}:{year}", o["value"], jsround(100 * ind_in[year] / dev_total[year], 2))
    elif aid in ("extfin.fdi.dev.greenfield_announced.IN.usd", "extfin.fdi.dev.greenfield_outward.IN.usd", "extfin.fdi.dev.mna_sales.IN.usd"):
        num = 14 if "greenfield_announced" in aid else 13 if "greenfield_outward" in aid else 5
        source = annual_annex(num, "India")
        for o in observations:
            year = int(o["date"][:4])
            same(f"{aid}:{year}", o["value"], jsround(source[year], 2))
    elif aid == "extfin.fdi.dev.inward_flow_recorded.IN.usd":
        for o in observations:
            year = int(o["date"][:4])
            same(f"{aid}:{year}", o["value"], jsround(ind_in[year], 2))
    elif aid == "extfin.fdi.dev.cumulative_per_person.usd":
        for row in table_rows:
            econ = "Republic of Korea" if row["economy"] == "Korea" else row["economy"]
            total_million = sum(raw(y, econ) or 0 for y in range(1990, 2026))
            same(f"{aid}:{econ}:cumulative_fdi_usd_billion", row["cumulative_fdi_usd_billion"], jsround(total_million / 1000))
            same(f"{aid}:{econ}:population_2025_million", row["population_2025_million"], jsround(population[econ] / 1e6))
            same(f"{aid}:{econ}:value", row["value"], jsround(total_million * 1e6 / population[econ]))
    elif aid == "extfin.fdi.dev.era_summary.usd":
        for row in table_rows:
            a, b = map(int, row["era"].split("-"))
            inward = sum(ind_in[y] for y in range(a, b + 1)) / 1000
            outward = sum(ind_out[y] for y in range(a, b + 1)) / 1000
            mean = sum(raw(y, "India", GFCF) for y in range(a, b + 1)) / (b - a + 1)
            same(f"{aid}:{a}-{b}:in", row["received_usd_billion"], jsround(inward, 2))
            same(f"{aid}:{a}-{b}:out", row["sent_out_usd_billion"], jsround(outward))
            same(f"{aid}:{a}-{b}:net", row["net_usd_billion"], jsround(inward - outward))
            same(f"{aid}:{a}-{b}:gfcf", row["avg_pct_of_capital_formation"], jsround(mean, 2))
    elif aid == "extfin.fdi.dev.inward_stock_gdp_peers.pct":
        for row in table_rows:
            same(f"{aid}:{row['economy']}", row["value"], jsround(raw(2025, row["economy"], GDP, "Stock"), 2))
        same(f"{aid}:row count", len(table_rows), len(stock_peers))
    elif aid == "extfin.fdi.dev.top100_mne_home_economy.count":
        dev, world = mne_counts(20), mne_counts(19)
        for row in table_rows:
            same(f"{aid}:{row['economy']}:developing", row["value"], dev[row["economy"]])
            same(f"{aid}:{row['economy']}:world", row["world_top100"], world[row["economy"]])
    else:
        errors.append({"artifact_without_rule": aid})

print(json.dumps({"checked_cells": checked, "mismatches": errors, "source_gaps": unverified}, indent=2))
raise SystemExit(1 if errors else 0)
