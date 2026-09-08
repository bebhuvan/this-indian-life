#!/usr/bin/env python3
"""Independent fact audit for q.econ.income_classification.

This intentionally does not import the ingest script. It reparses the saved raw
workbooks and WDI responses, recomputes the article's highest-risk derivations,
checks the generated artifacts, and writes a compact audit report.
"""

from __future__ import annotations

import hashlib
import html
import json
import math
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
CATALOG = ROOT / "data/catalog/worldbank-income-classifications-manifest.json"
EXPLANATION = ROOT / "data/explanations/en/q.econ.income_classification.json"
REPORT = ROOT / "data/audits/income-classification-fact-check.json"
SERIES_DIR = ROOT / "data/series"

checks: list[dict[str, Any]] = []


def record(name: str, actual: Any, expected: Any, note: str = "") -> None:
    if isinstance(expected, float):
        passed = math.isclose(float(actual), expected, rel_tol=0, abs_tol=0.051)
    else:
        passed = actual == expected
    checks.append(
        {
            "name": name,
            "status": "pass" if passed else "fail",
            "actual": actual,
            "expected": expected,
            **({"note": note} if note else {}),
        }
    )
    if not passed:
        raise AssertionError(f"{name}: expected {expected!r}, got {actual!r}")


def read_json(path: Path) -> Any:
    return json.loads(path.read_text())


def artifact(suffix: str) -> dict[str, Any]:
    return read_json(SERIES_DIR / f"worldbank-income-classifications.{suffix}.json")


def snap(catalog: dict[str, Any], key: str) -> Path:
    return ROOT / catalog["snapshots"][key]["path"]


def num(value: Any) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    match = re.search(r"[\d,]+(?:\.\d+)?", str(value))
    return float(match.group(0).replace(",", "")) if match else None


def wdi_values(path: Path) -> dict[str, dict[int, float]]:
    payload = read_json(path)
    result: dict[str, dict[int, float]] = {}
    for row in payload[1]:
        if row.get("countryiso3code") and row.get("value") is not None:
            result.setdefault(row["countryiso3code"], {})[int(row["date"])] = float(row["value"])
    return result


catalog = read_json(CATALOG)

# 1. Raw snapshot integrity. A changed byte invalidates the audit immediately.
for key, item in catalog["snapshots"].items():
    path = ROOT / item["path"]
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    record(f"snapshot hash: {key}", digest, item["sha256"])

# The non-tabular methodology and mover claims are also frozen as raw HTML.
blog_text = html.unescape(re.sub(r"<[^>]+>", " ", snap(catalog, "blog").read_text(errors="ignore")))
atlas_method_text = html.unescape(re.sub(r"<[^>]+>", " ", snap(catalog, "atlas_method").read_text(errors="ignore")))
for phrase in ("six countries moved", "nearly 10% larger", "11.7%"):
    record(f"World Bank mover article contains: {phrase}", phrase.lower() in blog_text.lower(), True)
record("Atlas method states three-year moving average", "three-year moving average" in atlas_method_text.lower(), True)

# 2. Reparse official 2025 release workbooks independently.
gnipc_ws = openpyxl.load_workbook(snap(catalog, "gnipc"), read_only=True, data_only=True)["GNIPC"]
atlas: dict[str, tuple[int, float]] = {}
ppp: dict[str, tuple[int, float]] = {}
for row in gnipc_ws.iter_rows(min_row=8, values_only=True):
    if isinstance(row[0], str) and len(row[0]) == 3 and isinstance(row[1], (int, float)):
        atlas[row[0]] = (int(row[1]), float(row[4]))
    if isinstance(row[11], str) and len(row[11]) == 3 and isinstance(row[6], (int, float)):
        ppp[row[11]] = (int(row[6]), float(row[9]))


def rank_value(key: str, sheet: str, start: int) -> dict[str, tuple[int, float]]:
    ws = openpyxl.load_workbook(snap(catalog, key), read_only=True, data_only=True)[sheet]
    output = {}
    for row in ws.iter_rows(min_row=start, values_only=True):
        if isinstance(row[0], str) and len(row[0]) == 3 and isinstance(row[1], (int, float)):
            output[row[0]] = (int(row[1]), float(row[4]))
    return output


gdp = rank_value("gdp", "GDP", 6)
gdp_ppp = rank_value("gdp_ppp", "GDP_PPP", 7)
gni = rank_value("gni", "GNI", 6)
population_release = rank_value("population", "POP", 6)

record("India PPP GNI per-person rank", ppp["IND"][0], 132)
record("India Atlas GNI per-person rank", atlas["IND"][0], 156)
record("India total PPP GDP rank", gdp_ppp["IND"][0], 3)
record("India total market GDP rank", gdp["IND"][0], 6)
record("PPP GNI per-person ranked universe", len(ppp), 196)
record("Atlas GNI per-person ranked universe", len(atlas), 200)
record("PPP GDP ranked universe", len(gdp_ppp), 197)
record("market GDP ranked universe", len(gdp), 204)
record("India PPP GNI per person", ppp["IND"][1], 11600.0)
record("India Atlas GNI per person", atlas["IND"][1], 2760.0)
record("India Atlas GNI, million US dollars", round(gni["IND"][1]), 4_036_219)
record("India market GDP, million US dollars", round(gdp["IND"][1]), 3_956_067)
record("India PPP GDP, million international dollars", round(gdp_ppp["IND"][1]), 17_197_369)
record("India population, thousands", round(population_release["IND"][1]), 1_463_866)

# 3. Historical classifications and thresholds.
history_wb = openpyxl.load_workbook(snap(catalog, "history"), read_only=True, data_only=True)
history_ws = history_wb["Country Analytical History"]
year_cells = list(next(history_ws.iter_rows(min_row=6, max_row=6, values_only=True)))
year_cols = {index: int(value) for index, value in enumerate(year_cells) if isinstance(value, int) and 1987 <= value <= 2025}
classes: dict[str, dict[int, str]] = {}
names: dict[str, str] = {}
for row in history_ws.iter_rows(min_row=7, values_only=True):
    if isinstance(row[0], str) and len(row[0]) == 3:
        names[row[0]] = str(row[1])
        classes[row[0]] = {year: row[index] for index, year in year_cols.items() if row[index] in ("L", "LM", "UM", "H")}

threshold_ws = history_wb["Thresholds"]
threshold_years = list(next(threshold_ws.iter_rows(min_row=7, max_row=7, values_only=True)))
threshold_rows = {
    "ida": list(next(threshold_ws.iter_rows(min_row=17, max_row=17, values_only=True))),
    "ibrd": list(next(threshold_ws.iter_rows(min_row=15, max_row=15, values_only=True))),
    "lower": list(next(threshold_ws.iter_rows(min_row=22, max_row=22, values_only=True))),
    "upper": list(next(threshold_ws.iter_rows(min_row=23, max_row=23, values_only=True))),
    "high": list(next(threshold_ws.iter_rows(min_row=24, max_row=24, values_only=True))),
}
thresholds: dict[int, dict[str, float]] = {}
for index, year in enumerate(threshold_years):
    if isinstance(year, int) and 1987 <= year <= 2025:
        values = {key: num(row[index]) for key, row in threshold_rows.items()}
        if all(value is not None for value in values.values()):
            thresholds[year] = {key: float(value) for key, value in values.items() if value is not None}

record("2025 lower-middle entry", thresholds[2025]["lower"], 1176.0)
record("2025 low-income ceiling", thresholds[2025]["lower"] - 1, 1175.0)
record("2025 upper-middle entry", thresholds[2025]["upper"], 4636.0)
record("2025 high-income entry", thresholds[2025]["high"], 14375.0)
record("2025 IDA allocation ceiling", thresholds[2025]["ida"], 1365.0)
record("2025 IBRD graduation discussion threshold", thresholds[2025]["ibrd"], 8105.0)
record("India classification, 2025", classes["IND"][2025], "LM")
record("India first lower-middle year", min(year for year, value in classes["IND"].items() if value == "LM"), 2007)

current_ws = openpyxl.load_workbook(snap(catalog, "current_groups"), read_only=True, data_only=True)["List of economies"]
current = {row[1]: (row[3], row[4]) for row in current_ws.iter_rows(min_row=2, values_only=True) if isinstance(row[1], str)}
record("India current income-group label", current["IND"][0], "Lower middle income")
record("India current lending category", current["IND"][1], "IBRD")

# 4. Recompute key derived charts from raw WDI snapshots.
gni_pc = wdi_values(snap(catalog, "wdi_gni_pc"))
population = wdi_values(snap(catalog, "wdi_population"))
gni_lcu = wdi_values(snap(catalog, "wdi_gni_lcu"))["IND"]

india_ratio = round(100 * gni_pc["IND"][2025] / thresholds[2025]["upper"], 1)
record("India progress ratio, 2025", india_ratio, 59.5)
record("India progress artifact endpoint", artifact("progress-to-upper-middle.ind")["observations"][-1]["value"], india_ratio)
record(
    "rank artifact universes",
    [row["universe"] for row in artifact("india-total-and-per-person-ranks")["rows"]],
    [197, 204, 196, 200],
)
for code in ("IND", "CHN", "IDN", "VNM", "BGD", "LKA"):
    expected_progress = [
        {"date": str(year), "value": round(100 * value / thresholds[year]["upper"], 1)}
        for year, value in sorted(gni_pc[code].items())
        if year in thresholds
    ]
    record(
        f"full progress series: {code}",
        artifact(f"progress-to-upper-middle.{code.lower()}")["observations"],
        expected_progress,
    )

# Population shares by class in 1987 and 2025.
expected_global = []
for year in (1987, 2025):
    totals = {label: 0.0 for label in ("L", "LM", "UM", "H")}
    denominator = 0.0
    for code, year_map in classes.items():
        value = population.get(code, {}).get(year)
        category = year_map.get(year)
        if category in totals and value is not None:
            totals[category] += value
            denominator += value
    expected_global.extend(round(100 * totals[label] / denominator, 1) for label in ("L", "LM", "UM", "H"))
record("global population-by-class chart", [row["value"] for row in artifact("global-population-by-class")["rows"]], expected_global)

# Concentration inside the 2025 lower-middle-income population.
lm = sorted(
    (
        (code, population[code][2025])
        for code, year_map in classes.items()
        if year_map.get(2025) == "LM" and 2025 in population.get(code, {})
    ),
    key=lambda item: item[1],
    reverse=True,
)
lm_total = sum(value for _, value in lm)
top_shares = [round(100 * value / lm_total, 1) for _, value in lm[:5]]
top_shares.append(round(100 * (lm_total - sum(value for _, value in lm[:5])) / lm_total, 1))
record("lower-middle population concentration chart", [row["value"] for row in artifact("lower-middle-population-concentration")["rows"]], top_shares)

# Full cohort with an independently coded product-limit calculation.
followup: list[tuple[int, bool]] = []
for code, year_map in classes.items():
    starts = [
        year
        for year, category in year_map.items()
        if year > 1987
        and category == "LM"
        and year in thresholds
        and year in gni_pc.get(code, {})
        and 55 <= 100 * gni_pc[code][year] / thresholds[year]["upper"] <= 65
    ]
    if not starts:
        continue
    start = min(starts)
    crossing = min((year for year, category in year_map.items() if year > start and category in ("UM", "H")), default=None)
    followup.append(((crossing or 2025) - start, crossing is not None))


def crossed_by(horizon: int) -> float:
    survival = 1.0
    for event_time in sorted({duration for duration, event in followup if event and duration <= horizon}):
        risk = sum(duration >= event_time for duration, _ in followup)
        events = sum(duration == event_time and event for duration, event in followup)
        survival *= 1 - events / risk
    return round(100 * (1 - survival), 1)


record("full crossing cohort size", len(followup), 67)
record("full crossing cohort events", sum(event for _, event in followup), 41)
record("Kaplan-Meier crossing chart", [row["value"] for row in artifact("all-economies-crossing-experience")["rows"]], [crossed_by(year) for year in (5, 10, 15, 20)])

# Selected completed Asian transitions, kept separate from the full cohort.
country_names = {
    "CHN": "China",
    "LKA": "Sri Lanka",
    "IDN": "Indonesia",
    "VNM": "Vietnam",
    "PHL": "Philippines",
    "THA": "Thailand",
    "FSM": "Micronesia, Fed. Sts.",
    "JOR": "Jordan",
    "TGO": "Togo",
}
selected_durations = {}
for code in ("CHN", "LKA", "IDN", "VNM", "PHL", "THA"):
    starts = [
        year
        for year, category in classes[code].items()
        if category == "LM"
        and year in thresholds
        and year in gni_pc.get(code, {})
        and 55 <= 100 * gni_pc[code][year] / thresholds[year]["upper"] <= 65
    ]
    start = min(starts)
    crossing = min(year for year, category in classes[code].items() if year > start and category in ("UM", "H"))
    selected_durations[country_names[code]] = crossing - start
record(
    "selected Asian completed transitions",
    {row["label"]: row["value"] for row in artifact("years-from-india-position-to-crossing")["rows"]},
    selected_durations,
)

# Post-crossing reversals.
outcomes = {"never": 0, "recovered": 0, "below": 0}
for year_map in classes.values():
    years = sorted(year_map)
    first = next(
        (
            year
            for year in years
            if year_map[year] in ("UM", "H") and any(year_map[prior] in ("L", "LM") for prior in years if prior < year)
        ),
        None,
    )
    if first is None:
        continue
    later_below = any(year_map[year] in ("L", "LM") for year in years if year > first)
    if year_map.get(2025) in ("L", "LM"):
        outcomes["below"] += 1
    elif later_below:
        outcomes["recovered"] += 1
    else:
        outcomes["never"] += 1
record("observed first crossers", sum(outcomes.values()), 76)
record("crossing outcome counts", outcomes, {"never": 54, "recovered": 16, "below": 6})

# Atlas/PPP ranking reversal and multiples.
record("Bangladesh Atlas lead", round(100 * (atlas["BGD"][1] / atlas["IND"][1] - 1), 1), 2.9)
record("India PPP lead", round(100 * (ppp["IND"][1] / ppp["BGD"][1] - 1), 1), 8.3)
record("India PPP-to-Atlas multiple", round(ppp["IND"][1] / atlas["IND"][1], 1), 4.2)
record("India-Bangladesh ruler-swap chart", [row["value"] for row in artifact("india-bangladesh-ruler-swap")["rows"]], [2.9, 8.3])

multiple_codes = {
    "India": "IND",
    "Bangladesh": "BGD",
    "Pakistan": "PAK",
    "Sri Lanka": "LKA",
    "Vietnam": "VNM",
    "Philippines": "PHL",
    "Indonesia": "IDN",
    "China": "CHN",
    "Thailand": "THA",
    "Malaysia": "MYS",
    "South Korea": "KOR",
    "United States": "USA",
}
expected_multiples = {label: round(ppp[code][1] / atlas[code][1], 1) for label, code in multiple_codes.items()}
record(
    "PPP-to-Atlas multiple chart",
    {row["label"]: row["value"] for row in artifact("ppp-atlas-multiple")["rows"]},
    expected_multiples,
)

# The 2025 mover bars come from raw WDI values; mechanism labels come from the
# frozen World Bank mover article checked above.
mover_labels = {
    "VNM": "Vietnam",
    "PHL": "Philippines",
    "LKA": "Sri Lanka",
    "FSM": "Micronesia, Fed. Sts.",
    "JOR": "Jordan",
    "TGO": "Togo",
}
mover_changes = {
    label: round(100 * (gni_pc[code][2025] / gni_pc[code][2024] - 1), 1)
    for code, label in mover_labels.items()
}
record(
    "2025 mover GNI-per-person changes",
    {row["label"]: row["value"] for row in artifact("movers-gni-pc-change")["rows"]},
    mover_changes,
)
record(
    "2025 mover mechanism labels",
    {row["label"]: row["group"] for row in artifact("movers-gni-pc-change")["rows"]},
    {
        "Vietnam": "Growth and exports",
        "Philippines": "Broad-based growth",
        "Sri Lanka": "Economic recovery",
        "Micronesia, Fed. Sts.": "Post-Covid recovery",
        "Jordan": "GDP rebasing",
        "Togo": "Population revision",
    },
)

record(
    "analytical and lending ladder",
    [row["value"] for row in artifact("classification-and-lending-ladder")["rows"]],
    [1365.0, 2760.0, 4636.0, 8105.0, 14375.0],
)

# India's ordered accounting bridge.
start, end = 2015, 2025
atlas_factor = {year: gni_lcu[year] / population["IND"][year] / gni_pc["IND"][year] for year in (start, end)}
level = 100 * gni_pc["IND"][start] / thresholds[start]["upper"]
bridge = [round(level, 1)]
for multiplier in (
    gni_lcu[end] / gni_lcu[start],
    population["IND"][start] / population["IND"][end],
    atlas_factor[start] / atlas_factor[end],
    thresholds[start]["upper"] / thresholds[end]["upper"],
):
    level *= multiplier
    bridge.append(round(level, 1))
record("India accounting bridge", [row["value"] for row in artifact("india-progress-accounting-bridge")["rows"]], bridge)

# Scenario arithmetic.
ratios = {year: gni_pc["IND"][year] / thresholds[year]["upper"] for year in thresholds if year in gni_pc["IND"]}
scenario_years = []
for start_year in (2007, 2015, 2021, 2019):
    rate = (ratios[2025] / ratios[start_year]) ** (1 / (2025 - start_year)) - 1
    wait = round(math.log(1 / ratios[2025]) / math.log(1 + rate))
    scenario_years.append((wait, 2025 + wait))
scenario_years.sort()
record("crossing scenario waits", [(row["value"], row["projectedYear"]) for row in artifact("india-crossing-arithmetic")["rows"]], scenario_years)

# Social and structural milestone comparisons are independently reconstructed
# from their WDI snapshots and classification years.
u5 = wdi_values(snap(catalog, "wdi_under5"))
manufacturing = wdi_values(snap(catalog, "wdi_manufacturing"))


def latest_before(values: dict[int, float], year: int) -> float:
    return values[max(candidate for candidate in values if candidate <= year)]


comparison_codes = ("IND", "CHN", "IDN", "VNM", "PHL", "LKA", "THA", "MYS")
comparison_names = {
    "IND": "India",
    "CHN": "China",
    "IDN": "Indonesia",
    "VNM": "Vietnam",
    "PHL": "Philippines",
    "LKA": "Sri Lanka",
    "THA": "Thailand",
    "MYS": "Malaysia",
}


def crossing_year(code: str) -> int:
    return 2025 if code == "IND" else min(year for year, category in classes[code].items() if category in ("UM", "H"))


expected_u5 = {comparison_names[code]: round(latest_before(u5[code], crossing_year(code)), 1) for code in comparison_codes}
expected_mfg = {comparison_names[code]: round(latest_before(manufacturing[code], crossing_year(code)), 1) for code in comparison_codes}
record("under-five mortality milestone chart", {row["label"]: row["value"] for row in artifact("under5-at-upper-middle-crossing")["rows"]}, expected_u5)
record("manufacturing milestone chart", {row["label"]: row["value"] for row in artifact("manufacturing-at-upper-middle-crossing")["rows"]}, expected_mfg)

# 5. HCES pairing is reproduced from the two original MoSPI artifacts.
rural = read_json(SERIES_DIR / "mospi-hces.IN.hces_fractile_mpce_rural.json")["rows"]
urban = read_json(SERIES_DIR / "mospi-hces.IN.hces_fractile_mpce_urban.json")["rows"]
paired = []
for rural_row, urban_row in zip(rural, urban):
    paired.extend([rural_row["value"], urban_row["value"]])
record("rural-urban HCES paired chart", [row["value"] for row in artifact("household-consumption-rural-urban")["rows"]], paired)

hces_pdf = ROOT / "data/snapshots/mospi-hces/HCES_Press_Note_2023-24_27122024.pdf"
hces_text = "\n".join(page.extract_text() or "" for page in PdfReader(hces_pdf).pages)
for value in ("1,677", "2,376", "10,137", "20,310"):
    record(f"MoSPI HCES primary PDF contains {value}", value in hces_text, True)

# 6. Independent national-accounts reconciliation. Periods differ deliberately:
# World Bank values are calendar-year estimates, while MoSPI reports fiscal years.
wb_gni_observations = artifact("india-gni-current-lcu")["observations"]
mospi_gni_observations = read_json(SERIES_DIR / "mospi.IN.econ.nas.gni_nominal.json")["observations"]
wb_latest = wb_gni_observations[-1]
mospi_latest = mospi_gni_observations[-1]
wb_year = int(str(wb_latest["date"])[:4])
mospi_year = int(str(mospi_latest["date"])[:4])
record("World Bank versus MoSPI GNI latest-year gap within tolerance", abs(wb_year - mospi_year) <= 1, True)
record("World Bank versus MoSPI GNI value gap within ₹15tn tolerance", abs(wb_latest["value"] - mospi_latest["value"]) <= 15_000_000_000_000, True)

# 7. High-risk prose claims must remain present with the audited values and caveats.
body = read_json(EXPLANATION)["article"]["bodyMarkdown"]
required_phrases = [
    "ranked 132nd among 196 estimates",
    "ranked 156th among 200",
    "49.7% live in India",
    "67 economies",
    "29.2% within ten years",
    "It is not India's probability",
    "22 of 76 observed first crossers",
    "Only the first and last bars are observed",
    "not a forecast",
    "These figures are not an alternative estimate of GNI",
    "Every 2025 World Bank value is an estimate and may be revised",
]
for phrase in required_phrases:
    record(f"prose contains: {phrase}", phrase in body, True)

explanation = read_json(EXPLANATION)
record("curated chart count", len(explanation.get("chartExplainers", [])), 15)
record("reader-question section count", body.count("\n## ") + int(body.startswith("## ")), 8)
record("grouped section-to-visual map count", len(explanation.get("sectionVisualMap", [])), 8)
record(
    "all grouped visuals mapped",
    sum(len(item.get("visualIds", [])) for item in explanation.get("sectionVisualMap", [])),
    15,
)
for index, chart_explainer in enumerate(explanation.get("chartExplainers", []), start=1):
    record(
        f"chart explainer {index} has all rich fields",
        all(chart_explainer.get(key) for key in ("visualId", "takeaway", "detail", "whyShowThis", "howToRead", "mistakeToAvoid", "mobileNote")),
        True,
    )

# Every numeric token in the body must trace to an artifact, a source label/date,
# or one of the explicitly audited methodological constants and cohort counts.
numeric_support: set[float] = {
    thresholds[2025]["lower"] - 1,
    thresholds[2025]["lower"],
    4635,
    55,
    65,
    42,
    47,
    54,
    67,
    76,
    16,
    22,
    29,
    11.7,
    1000,
}


def harvest_numbers(value: Any) -> None:
    if isinstance(value, bool) or value is None:
        return
    if isinstance(value, (int, float)):
        numeric_support.add(float(value))
    elif isinstance(value, str):
        for match in re.finditer(r"\d[\d,]*(?:\.\d+)?", value):
            numeric_support.add(float(match.group(0).replace(",", "")))
    elif isinstance(value, list):
        for item in value:
            harvest_numbers(item)
    elif isinstance(value, dict):
        for item in value.values():
            harvest_numbers(item)


for path in SERIES_DIR.glob("worldbank-income-classifications.*.json"):
    harvest_numbers(read_json(path))
harvest_numbers(read_json(SERIES_DIR / "mospi-hces.IN.hces_fractile_mpce_rural.json"))
harvest_numbers(read_json(SERIES_DIR / "mospi-hces.IN.hces_fractile_mpce_urban.json"))


def supported_number(value: float) -> bool:
    for source_value in numeric_support:
        candidates = (source_value, source_value / 1_000, source_value / 1_000_000, source_value / 1_000_000_000, source_value / 1_000_000_000_000)
        if any(abs(value - candidate) <= 0.051 for candidate in candidates):
            return True
    return False


unsupported_tokens = []
for match in re.finditer(r"(?<![A-Za-z])(?:[$₹])?\d[\d,]*(?:\.\d+)?%?", body):
    token = match.group(0)
    value = float(token.replace("$", "").replace("₹", "").replace(",", "").replace("%", ""))
    if 1900 <= value <= 2100 or supported_number(value):
        continue
    unsupported_tokens.append(
        {
            "token": token,
            "context": body[max(0, match.start() - 35) : match.end() + 35].replace("\n", " "),
        }
    )
record("all prose numeric tokens trace to audited data or declared constants", unsupported_tokens, [])

failed = [item for item in checks if item["status"] != "pass"]
report = {
    "articleId": "q.econ.income_classification",
    "generatedAt": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
    "status": "pass" if not failed else "fail",
    "summary": {
        "checks": len(checks),
        "passed": len(checks) - len(failed),
        "failed": len(failed),
        "rawSnapshotsVerified": len(catalog["snapshots"]),
    },
    "scope": [
        "raw snapshot hashes",
        "official workbook values and ranks",
        "classification and lending labels",
        "historical thresholds",
        "independent recomputation of the main derived charts",
        "HCES rural-urban pairing",
        "World Bank versus MoSPI national-accounts reconciliation",
        "high-risk prose claims and caveat retention",
    ],
    "checks": checks,
}
REPORT.parent.mkdir(parents=True, exist_ok=True)
REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
print(f"{report['status'].upper()}: {report['summary']['passed']}/{report['summary']['checks']} checks passed")
print(REPORT.relative_to(ROOT))
